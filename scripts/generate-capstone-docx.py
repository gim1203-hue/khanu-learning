from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Khanu Learning Capstone Documentation.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.72)
section.left_margin = Inches(.78)
section.right_margin = Inches(.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(9.5)
styles['Normal'].font.color.rgb = RGBColor(44, 50, 46)
for name, size, color in [('Title', 30, '4734A5'), ('Heading 1', 18, '4734A5'), ('Heading 2', 12, '24664D')]:
    styles[name].font.name = 'Aptos Display'
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string(color)

header = section.header.paragraphs[0]
header.text = 'KHANU LEARNING  /  PROJECT DOCUMENTATION'
header.style = 'Caption'
header.runs[0].font.color.rgb = RGBColor(100, 106, 101)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Khanu Learning  •  Capstone project  •  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def para(text='', style=None):
    return doc.add_paragraph(text, style)

def heading(text, level=1):
    doc.add_heading(text, level=level)

def bullets(items):
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Shading Accent 1'
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
    doc.add_paragraph()
    return t

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(74)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('KHANU')
r.bold = True
r.font.size = Pt(15)
r.font.color.rgb = RGBColor(36, 102, 77)
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('Learning Platform')
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Capstone project plan, solution, design, build, and deployment guide')
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(85, 92, 87)
para('A tutoring discovery and appointment-planning experience with an animated study buddy, goal-based matching, and two deployment modes.', None).alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

heading('1. Executive summary')
para('Khanu Learning is a responsive tutoring discovery and appointment-planning web application. It brings a sample service catalog, goal-based recommendations, booking workflow, and learner dashboard into one place. The Khanu robot provides a recognizable, customizable study-buddy identity.')
para('The full-stack application uses React and Vite, an Express REST API, MongoDB through Mongoose, JWT authentication, and an optional local language-model explanation. A GitHub Pages edition offers a static browser demo. The Pages demo can browse, match, and save sample bookings in that browser, but it is not secure account authentication and it does not share or persist records on a server.')

heading('2. Problem, audience, and solution')
heading('Problem', 2)
para('Learners may know they need help but struggle to choose a subject, tutor, learning level, price, or next step. Searching a list of disconnected offerings does not explain which option fits their goal or provide a simple way to organize a session.')
heading('Intended audience', 2)
bullets(['Middle- and high-school learners looking for subject support.', 'Test-prep and language learners comparing focused sessions.', 'College applicants who want writing and application guidance.', 'Parents or guardians helping a learner plan next steps.'])
heading('Solution', 2)
para('Khanu lets visitors filter a sample catalog, describe a learning goal, receive up to three ranked matches with plain-language reasons, select a future appointment, and see upcoming sessions in a dashboard. The matching logic considers goal words, subject, level, and budget. The interface is responsive and includes a theme switch and robot-color customization.')

heading('3. Project goals and scope')
bullets(['Provide a clear, accessible catalog and navigation on desktop and mobile.', 'Make session selection easier with goal-aware recommendations and reasons.', 'Support account creation, login, booking, ownership checks, cancellation, and MongoDB persistence in full-stack mode.', 'Provide an unmistakably labeled, password-free static browser demo for GitHub Pages.', 'Keep the deterministic matching feature usable without any language model.'])
para('Not in scope: payments, a real tutor marketplace, verified instructor accounts, live availability, calendar synchronization, messaging, and video lessons. Tutor names, prices, ratings, and time slots are example catalog content, not a live service commitment.')

heading('4. Included tutoring services')
table(['Sample service', 'Area / level', 'Duration / price'], [
    ('Algebra Foundations', 'Math · middle school', '60 minutes · $42'),
    ('SAT Math Boost', 'Test prep', '75 minutes · $58'),
    ('Essay Writing Workshop', 'Writing · all levels', '60 minutes · $45'),
    ('Chemistry, Made Clear', 'Science · high school', '60 minutes · $49'),
    ('Spanish Conversation Club', 'Languages · beginner/intermediate', '45 minutes · $36'),
    ('College Application Coaching', 'College prep', '60 minutes · $65'),
])

heading('5. User journeys')
heading('Discover and match', 2)
para('A learner browses or filters services, writes a goal, selects optional subject/level/budget details, and submits the form. Khanu returns a ranked short list and reasons. The ranker assigns higher weight to title matches, then contextual matches, level fit, and affordability.')
heading('Book and manage (full-stack)', 2)
para('A visitor registers or signs in. The backend verifies credentials, stores a password hash, and issues an expiring signed token. A signed-in learner chooses a future date/time, optionally enters a note, and sees the booking in My Sessions. Only that learner can list or cancel that record.')
heading('Try the Pages edition', 2)
para('A visitor enters a display name to create a browser-only demo profile. Khanu collects no email or password. The sample booking flow works in browser local storage, including after refresh, but the record is not synchronized and is not a real appointment.')

heading('6. Architecture and services')
table(['Layer', 'Technology', 'Responsibilities'], [
    ('Frontend', 'React 19, Vite, CSS', 'Catalog, search/filter, matching form, robot, account UI, booking modal, dashboard, theme.'),
    ('API', 'Node.js, Express 5, Zod', 'REST routes, validation, auth, booking rules, ownership, health status.'),
    ('Database', 'MongoDB, Mongoose', 'User, service, and booking documents; service seeding; persistence.'),
    ('Matching', 'Deterministic JavaScript ranker', 'Goal/subject/level/budget scoring with reason labels.'),
    ('Optional explanation', 'Local OpenAI-compatible model endpoint', 'Rephrases the ranked results when a local model is available; never required.'),
    ('Static host', 'GitHub Pages + GitHub Actions', 'Builds and publishes the static browser-demo frontend.'),
])
para('Full-stack request flow: React → JSON REST API → controllers/services → MongoDB and/or local model. Pages request flow: React → in-browser catalog, matcher, and localStorage adapter. GitHub Pages does not execute Express, MongoDB, or a visitor local model process.')

heading('7. Data model and API')
table(['Record', 'Fields and relationship'], [
    ('User', 'Name, unique normalized email, bcrypt password hash, timestamps. One user can own many bookings.'),
    ('Service', 'Unique slug, title, category, duration, price, level, format, tutor, rating, tags, description.'),
    ('Booking', 'Owner user ID, service slug/title/tutor snapshot, start time, note, status, timestamps.'),
])
table(['Method and path', 'Purpose'], [
    ('GET /api/health', 'Backend and database status.'),
    ('GET /api/services', 'Catalog listing with optional query/category filters.'),
    ('POST /api/auth/register', 'Register and issue a signed session token.'),
    ('POST /api/auth/login', 'Verify account credentials and issue a token.'),
    ('GET /api/auth/me', 'Read the signed-in user.'),
    ('POST /api/ai/recommend', 'Return up to three scored matches and explanations.'),
    ('GET/POST /api/bookings', "List current user's bookings / create a future booking."),
    ('PATCH /api/bookings/:id/cancel', 'Cancel a booking owned by the current user.'),
])

heading('8. Security and privacy')
bullets(['Backend passwords are hashed with bcrypt; raw passwords are not stored.', 'JWT signing secrets and MongoDB connection strings stay on the server.', 'Protected API routes authenticate users; booking queries are scoped to the signed-in owner.', 'Request data is validated and registration/login are rate-limited.', 'The local matcher processes goals on the same machine/browser. The optional model endpoint is local and configurable.', 'Pages demo mode uses local storage and is not secure authentication. It deliberately asks only for a display name; do not enter real student records or sensitive notes.', 'Any VITE_* setting is bundled in public frontend code. It must never contain a password, secret, or private token.'])

heading('9. Source layout')
para('Root files contain package scripts, lockfile, environment template, and setup instructions. frontend/src contains React screens, reusable components, styles, auth state, API adapter, static demo data, and local matcher. backend/src contains Express startup, API routes/controllers, database configuration, Mongoose models, auth middleware, data services, and tests. scripts contains the GitHub Pages build helper and this document generator. docs contains the detailed project guides. `.github/workflows/pages.yml` automates Pages build and publishing.')

heading('10. Build and validation')
bullets(['Run `npm ci` to install the locked JavaScript dependencies.', 'Run `npm run dev` for the local Vite frontend and Express API.', 'Run `npm test` for API auth/booking, backend matching, and Pages browser-demo checks.', 'Run `npm run build` to produce the frontend served by the Express app.', 'Run `npm run build:pages` to create the static Pages artifact with the repository path.', 'With MongoDB configured, verify that accounts, catalog seed records, and bookings survive restart. Without it, the local banner explains the in-memory demo reset behavior.'])
para('The project tests cover sample API registration, route protection, booking/list/cancel behavior, recommendation ranking, static demo matching, and browser-local bookings. A production release still requires an external hosting configuration and MongoDB integration test.')

heading('11. Open in VS Code and run locally')
bullets(['Open the `local-assistant` project folder in VS Code.', 'Install Node.js 20 or newer.', 'Open a terminal in the project folder and run `npm ci`.', 'Copy `.env.example` to `.env`; set a randomly generated `JWT_SECRET`. Set `MONGODB_URI` to a local/managed database for persistence.', 'Run `npm run dev`; open `http://127.0.0.1:5173`. For a production-like run use `npm run build` then `npm start` at port 4173.', 'On Windows, `Start Assistant.bat` opens the local production app. The optional model is not needed for recommendations.'])

heading('12. GitHub Pages deployment')
bullets(['Create a GitHub repository and push the project source on `main` or `master`.', 'In repository Settings → Pages, select GitHub Actions as the publishing source.', 'The included workflow installs dependencies, builds `frontend/dist`, uploads the artifact, and deploys it.', 'Without an API URL, Pages serves the labeled browser demo. Browser profile and appointment data remain on that browser.', 'For production accounts and shared persistent data, deploy Express and MongoDB separately, store secrets on that host, set a CORS allow-list for the Pages origin, and define the public `VITE_API_BASE_URL` repository variable. Trigger the Pages workflow again.', 'The local Qwen process cannot be reached by a hosted service. Keep deterministic matching or intentionally configure a secure server-side model endpoint.'])
para('No repository remote or public Pages URL is included in this local source package. Public deployment must be completed from the selected GitHub repository after choosing its name and visibility. See `docs/DEPLOYMENT.md` for the full checklist.')

heading('13. Included documentation')
bullets(['README.md: entry point, quick start, and current status.', 'docs/PROJECT.md: problem, audiences, requirements, user stories, and acceptance checks.', 'docs/DESIGN.md: wireframe, architecture, data model, and security.', 'docs/PROJECT_MAP.md: file-by-file source guide.', 'docs/USER_GUIDE.md: learner-facing workflows and mode differences.', 'docs/API_REFERENCE.md: endpoint methods and example request bodies.', 'docs/AI_AND_TESTING.md: ranking rationale, test matrix, and disclosure.', 'docs/BUILD_GUIDE.md: VS Code setup, commands, and troubleshooting.', 'docs/DEPLOYMENT.md: GitHub Pages and separate backend steps.', 'docs/SECURITY_AND_PRIVACY.md: data boundaries and product limitations.'])

heading('14. Completion and next configuration')
para('The local source package includes the React experience, Express backend, Mongo models, matching logic, test suite, Pages workflow, browser-only Pages demo, and project documentation. The local production build and automated tests should be run before each submission. A public repository, public site URL, production database, and production backend are not configured in this source copy; they require the owner to choose/authorize those destinations and provide their service settings through secret managers.')

doc.core_properties.title = 'Khanu Learning Capstone Project Documentation'
doc.core_properties.subject = 'Project problem, design, architecture, services, build, verification, and deployment'
doc.core_properties.author = 'Khanu Learning project team'
doc.save(OUT)
print(OUT)
